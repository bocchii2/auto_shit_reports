from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from app.services import storage
from app.services.csv_import import load_csv
from app.services.excel_generator import generate_excel
from app.services.filters import apply_filters
from app.services.report_code import (
    build_codigo,
    build_periodo_corto,
    build_periodo_label,
    format_fecha_informe,
)
from app.services.word_generator import generate_word


def main() -> None:
    # reset
    db = ROOT / "data" / "app.db"
    if db.exists():
        db.unlink()
    exp = ROOT / "data" / "exports"
    if exp.exists():
        for child in exp.iterdir():
            if child.is_dir():
                import shutil

                shutil.rmtree(child)
            else:
                child.unlink()

    storage.init_db()
    storage.create_profile(
        name="Jostin Cedeño",
        nombre="Jostin Abraham Cedeño Mosquera",
        cargo="Programador Junior Software",
        cargo_firma="Parametrizador Jr.",
        supervisor="Ing. Patricio Guzman",
        cargo_supervisor="Supervisor Inmediato",
        empresa="Altura S.A.",
        equipo="IDS",
        ciudad="Manta",
        is_default=True,
    )
    print("profiles", [p["name"] for p in storage.list_profiles()])

    rows, _ = load_csv(ROOT / "sample_data" / "actividades_ejemplo.csv")
    bid = storage.save_import_batch("actividades_ejemplo.csv", rows)
    acts = storage.get_activities(bid)
    jul = apply_filters(acts, mes=7, anio=2026)
    print("jul rows", len(jul))

    today = date.today()
    header = {
        "periodo_mes": 7,
        "periodo_anio": 2026,
        "codigo_informe": build_codigo(7, 2026),
        "nombre": "Jostin Abraham Cedeño Mosquera",
        "cargo": "Programador Junior Software",
        "cargo_firma": "Parametrizador Jr.",
        "supervisor": "Ing. Patricio Guzman",
        "cargo_supervisor": "Supervisor Inmediato",
        "empresa": "Altura S.A.",
        "equipo": "IDS",
        "ciudad": "Manta",
        "periodo": build_periodo_label(7, 2026),
        "periodo_corto": build_periodo_corto(7, 2026),
        "fecha_informe": format_fecha_informe(today.day, today.month, today.year),
    }
    out = ROOT / "data" / "exports" / header["codigo_informe"]
    out.mkdir(parents=True, exist_ok=True)
    w = generate_word(
        ROOT / "app" / "templates" / "informe.docx",
        out / f"{header['codigo_informe']}.docx",
        header=header,
        activities=jul,
    )
    x = generate_excel(
        ROOT / "app" / "templates" / "matriz.xlsx",
        out / f"Matriz-{header['codigo_informe']}.xlsx",
        header=header,
        activities=jul,
    )
    print("word bytes", w.stat().st_size, "excel bytes", x.stat().st_size)

    from openpyxl import load_workbook
    from zipfile import ZipFile
    from xml.etree import ElementTree as ET

    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    with ZipFile(w) as z:
        h = ET.fromstring(z.read("word/header1.xml"))
        print("HEADER:", [t.text for t in h.iter(ns + "t") if t.text])
    wb = load_workbook(x)
    ws = wb.active
    print("EXCEL nombre", ws["E3"].value, "periodo", ws["E5"].value)
    print("FINAL OK")


if __name__ == "__main__":
    main()
