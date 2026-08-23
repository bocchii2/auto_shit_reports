import sys
from docx import Document
from docx.oxml.ns import qn

doc = Document("data/exports/test_word.docx")
print("python-docx tables:", len(doc.tables))
for i, t in enumerate(doc.tables):
    all_text = " ".join(c.text or "" for r in t.rows for c in r.cells)
    first_cell = t.rows[0].cells[0].text if t.rows else ""
    nos = [r.cells[0].text for r in t.rows if r.cells]
    print(f"--- python-docx table[{i}] rows={len(t.rows)} first={first_cell!r}")
    if "actividad" in all_text.lower():
        print("  ACTIVITIES TABLE nos:", nos)

# SDT search
sdt_tables = []
for child in doc.element.body:
    if child.tag == qn("w:sdt"):
        tbls = child.findall(qn("w:tbl"))
        sdt_tables.append(len(tbls))
print("sdt blocks with tables:", sdt_tables)
for sdt in doc.element.body:
    if sdt.tag == qn("w:sdt"):
        tbls = sdt.findall(qn("w:tbl"))
        # also check nested sdtContent
        for sc in sdt.iter(qn("w:sdtContent")):
            tbls.extend(sc.findall(qn("w:tbl")))
        joined_all = "".join(t.text or "" for t in sdt.iter(qn("w:t"))).lower()
        if "actividad" in joined_all:
            tbl = sdt.iter(qn("w:tbl"))
            tbl = list(sdt.iter(qn("w:tbl")))[0] if "actividad" in joined_all else None
            rows = list(tbl.iterchildren(qn("w:tr")))
        nos = []
        for tr in rows[1:]:
            seen, row_nos = set(), []
            for tc in tr.iter(qn("w:tc")):
                if id(tc) in seen:
                    continue
                seen.add(id(tc))
                row_nos.append("".join(x.text for x in tc.iter(qn("w:t"))))
            nos.append(row_nos[0] if row_nos else "")
        print("SDT activities nos:", nos)
